from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from app.resume_intel import extract_job_keywords


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return slug or "job"


def choose_relevant_lines(resume_lines: list[str], job_text: str, limit: int = 5) -> list[str]:
    keywords = extract_job_keywords(job_text)
    if not keywords:
        return resume_lines[:limit]

    scored: list[tuple[int, str]] = []
    lowered_keywords = [keyword.lower() for keyword in keywords]
    for line in resume_lines:
        score = sum(1 for keyword in lowered_keywords if keyword in line.lower())
        scored.append((score, line))

    ranked = [line for score, line in sorted(scored, key=lambda item: item[0], reverse=True) if line.strip()]
    return ranked[:limit]


def build_summary(job: dict, profile: dict) -> str:
    title = job.get("title") or "this role"
    company = job.get("company") or "your team"
    matched = extract_job_keywords((job.get("raw_text") or "") + " " + (job.get("description") or ""))
    resume_skills = profile.get("skills", [])
    aligned = [skill for skill in resume_skills if skill in matched] or resume_skills[:4]

    if aligned:
        core = ", ".join(aligned[:4])
        return (
            f"Systems-oriented operator with hands-on experience across {core}. "
            f"Targeting the {title} opportunity at {company} with a focus on shipping reliable workflows, "
            f"clean execution, and measurable operational lift."
        )

    return (
        f"Systems-oriented operator targeting the {title} opportunity at {company}, "
        f"with a focus on execution, automation, and operational ownership."
    )


def build_cover_letter(job: dict, profile: dict, relevant_lines: list[str], candidate_name: str) -> str:
    title = job.get("title") or "the role"
    company = job.get("company") or "your team"
    location = job.get("location") or "Remote"
    matched = extract_job_keywords((job.get("raw_text") or "") + " " + (job.get("description") or ""))
    matched_text = ", ".join(matched[:5]) if matched else "execution, systems thinking, and automation"
    evidence = relevant_lines[:2]

    body = [
        "Dear Hiring Team,",
        "",
        f"I'm applying for the {title} role at {company}. What stands out to me is the mix of ownership, execution, and systems work required to perform well in this seat.",
        "",
        f"My background aligns most closely with {matched_text}. I tend to work as an operator who closes the loop end to end: building the workflow, tightening the process, and making sure the thing actually runs in production instead of living as a half-finished concept.",
        "",
    ]

    if evidence:
        body.append("A few signals from my background that are relevant here:")
        for line in evidence:
            body.append(f"- {line}")
        body.append("")

    body.extend(
        [
            f"If selected, I would bring urgency, high ownership, and a bias toward finished systems over pretty plans. I'm interested in the {title} opportunity in {location} because it looks like a role where execution matters.",
            "",
            "Thank you for your time and consideration.",
            "",
            candidate_name,
        ]
    )

    return "\n".join(body)


def write_cover_letter(output_dir: Path, filename_prefix: str, cover_letter: str) -> Path:
    output_path = output_dir / f"{filename_prefix}_cover_letter.txt"
    output_path.write_text(cover_letter, encoding="utf-8")
    return output_path


def write_tailored_resume_docx(
    output_dir: Path,
    filename_prefix: str,
    candidate_name: str,
    candidate_email: str,
    candidate_phone: str,
    candidate_city: str,
    job: dict,
    summary: str,
    relevant_lines: list[str],
    base_resume_text: str,
) -> Path:
    doc = Document()
    doc.add_heading(candidate_name, level=0)

    contact_line = " | ".join([value for value in [candidate_email, candidate_phone, candidate_city] if value])
    if contact_line:
        doc.add_paragraph(contact_line)

    doc.add_heading("Target Role", level=1)
    doc.add_paragraph(f"{job.get('title', 'Role')} — {job.get('company', 'Company')}")

    doc.add_heading("Tailored Summary", level=1)
    doc.add_paragraph(summary)

    doc.add_heading("Most Relevant Evidence", level=1)
    for line in relevant_lines:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Base Resume Content", level=1)
    for raw_line in base_resume_text.splitlines():
        line = raw_line.strip()
        if line:
            doc.add_paragraph(line)

    output_path = output_dir / f"{filename_prefix}_tailored_resume.docx"
    doc.save(output_path)
    return output_path


def create_packet(settings, job: dict, profile: dict) -> dict[str, str]:
    filename_prefix = f"job_{job['id']}_{slugify(job['company'] or 'company')}_{slugify(job['title'])}"
    base_resume_text = profile["resume_text"]
    evidence_lines = profile["evidence_lines"]
    relevant_lines = choose_relevant_lines(evidence_lines, (job.get("raw_text") or "") + " " + (job.get("description") or ""))
    summary = build_summary(job, profile)
    cover_letter = build_cover_letter(job, profile, relevant_lines, settings.candidate_name)

    resume_output = write_tailored_resume_docx(
        output_dir=settings.output_dir,
        filename_prefix=filename_prefix,
        candidate_name=settings.candidate_name,
        candidate_email=settings.candidate_email,
        candidate_phone=settings.candidate_phone,
        candidate_city=settings.candidate_city,
        job=job,
        summary=summary,
        relevant_lines=relevant_lines,
        base_resume_text=base_resume_text,
    )
    cover_output = write_cover_letter(settings.output_dir, filename_prefix, cover_letter)

    return {
        "resume_output_path": str(resume_output),
        "cover_letter_path": str(cover_output),
        "cover_letter_text": cover_letter,
    }